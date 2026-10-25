"""
generate_master_skripsi.py
Menggabungkan seluruh bagian skripsi menjadi Master Dokumen Lengkap:
  1. Front Matter (Sampul Luar, Sampul Dalam, Pernyataan Keaslian, Persetujuan, Pengesahan,
     Persembahan, Motto, Kata Pengantar, Abstrak ID, Abstract EN, Daftar Isi, Daftar Tabel,
     Daftar Gambar, Daftar Lampiran) -> Penomoran Romawi (i, ii, ...)
  2. BAB I : PENDAHULUAN -> Penomoran Arab (1, 2, ...)
  3. BAB II: LANDASAN TEORI
  4. BAB III: METODE PENELITIAN
  5. BAB IV: HASIL PENELITIAN DAN PEMBAHASAN
  6. BAB V : PENUTUP
  7. Back Matter (DAFTAR PUSTAKA 32 Referensi APA 7th & LAMPIRAN 1-3)
Hasil akhir: 03_Naskah_Skripsi/SKRIPSI_FULL_ZAINI_LEON_MUSTHOFA_KAMAL.docx
"""

import os
import sys
import docx
from docx import Document
from docxcompose.composer import Composer
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

# Tambahkan direktori generators ke sys.path
generators_dir = os.path.dirname(os.path.abspath(__file__))
if generators_dir not in sys.path:
    sys.path.insert(0, generators_dir)

def assemble_full_skripsi():
    # 1. Regenerate all components to ensure fresh state
    print("[1/4] Memastikan seluruh komponen naskah terkompilasi...")
    import generate_front_matter
    generate_front_matter.generate_front_matter()
    
    import generate_back_matter
    generate_back_matter.generate_back_matter()

    workspace_root = os.path.abspath(os.path.join(generators_dir, "..", ".."))
    front_matter_path = os.path.join(workspace_root, "05_Source_Code", "front_matter.docx")
    bab1_path = os.path.join(workspace_root, "03_Naskah_Skripsi", "BAB_1_PENDAHULUAN_ZAINI_LEON.docx")
    bab2_path = os.path.join(workspace_root, "03_Naskah_Skripsi", "BAB_2_LANDASAN_TEORI_ZAINI_LEON.docx")
    bab3_path = os.path.join(workspace_root, "03_Naskah_Skripsi", "BAB_3_METODE_PENELITIAN_ZAINI_LEON.docx")
    bab4_path = os.path.join(workspace_root, "03_Naskah_Skripsi", "BAB_4_HASIL_DAN_PEMBAHASAN_ZAINI_LEON.docx")
    bab5_path = os.path.join(workspace_root, "03_Naskah_Skripsi", "BAB_5_PENUTUP_ZAINI_LEON.docx")
    back_matter_path = os.path.join(workspace_root, "05_Source_Code", "back_matter.docx")

    print("[2/4] Menginisialisasi Composer dokumen induk...")
    # Load Front Matter as master base
    master_doc = Document(front_matter_path)
    
    # Tambahkan Section Break untuk memisahkan Front Matter (Romawi) dan Isi Naskah (Arab)
    master_doc.add_section()
    
    composer = Composer(master_doc)

    # List of sequential chapter documents
    chapters = [
        ("BAB I", bab1_path),
        ("BAB II", bab2_path),
        ("BAB III", bab3_path),
        ("BAB IV", bab4_path),
        ("BAB V", bab5_path),
        ("DAFTAR PUSTAKA & LAMPIRAN", back_matter_path)
    ]

    print("[3/4] Menggabungkan bab-bab inti dan materi akhir secara berurutan...")
    for label, chap_path in chapters:
        if not os.path.exists(chap_path):
            raise FileNotFoundError(f"File {chap_path} tidak ditemukan!")
        sub_doc = Document(chap_path)
        composer.append(sub_doc)
        print(f"  -> Berhasil menggabungkan: {label} ({chap_path})")

    # Final destination path
    out_dir = os.path.join(workspace_root, "03_Naskah_Skripsi")
    os.makedirs(out_dir, exist_ok=True)
    out_master_path = os.path.join(out_dir, "SKRIPSI_FULL_ZAINI_LEON_MUSTHOFA_KAMAL.docx")
    composer.save(out_master_path)

    # Verifikasi dan konfigurasi penomoran halaman pada Section 1 (Romawi) dan Section 2 (Arab)
    doc_final = Document(out_master_path)
    
    # Section 0: Front Matter (i, ii, iii, ...)
    sec0 = doc_final.sections[0]
    sec0.top_margin = Cm(4.0)
    sec0.left_margin = Cm(4.0)
    sec0.bottom_margin = Cm(3.0)
    sec0.right_margin = Cm(3.0)
    sectPr0 = sec0._sectPr
    for el in list(sectPr0):
        if el.tag.endswith('pgNumType'):
            sectPr0.remove(el)
    pnt0 = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="lowerRoman" w:start="1"/>')
    sectPr0.append(pnt0)
    
    footer0 = sec0.footer
    p_ft0 = footer0.paragraphs[0]
    p_ft0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ft0.text = ""
    r0 = p_ft0.add_run()
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(10)
    p_ft0._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>'))

    # Section 1: Main Body (1, 2, 3, ...)
    if len(doc_final.sections) > 1:
        sec1 = doc_final.sections[1]
        sec1.top_margin = Cm(4.0)
        sec1.left_margin = Cm(4.0)
        sec1.bottom_margin = Cm(3.0)
        sec1.right_margin = Cm(3.0)
        sec1.footer.is_linked_to_previous = False
        sectPr1 = sec1._sectPr
        for el in list(sectPr1):
            if el.tag.endswith('pgNumType'):
                sectPr1.remove(el)
        pnt1 = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="decimal" w:start="1"/>')
        sectPr1.append(pnt1)

        footer1 = sec1.footer
        p_ft1 = footer1.paragraphs[0]
        p_ft1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ft1.text = ""
        r1 = p_ft1.add_run()
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(10)
        p_ft1._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>'))

    # Untuk section berikutnya (jika ada), teruskan format penomoran
    for s_idx in range(2, len(doc_final.sections)):
        sec_sub = doc_final.sections[s_idx]
        sec_sub.top_margin = Cm(4.0)
        sec_sub.left_margin = Cm(4.0)
        sec_sub.bottom_margin = Cm(3.0)
        sec_sub.right_margin = Cm(3.0)
        sec_sub.footer.is_linked_to_previous = True

    doc_final.save(out_master_path)
    file_size_mb = os.path.getsize(out_master_path) / (1024 * 1024)
    print(f"\n=======================================================")
    print(f" MASTER SKRIPSI LENGKAP SIAP!")
    print(f" File: {out_master_path}")
    print(f" Ukuran: {file_size_mb:.2f} MB")
    print(f" Jumlah Bagian / Sections: {len(doc_final.sections)}")
    print(f" Jumlah Paragraf: {len(doc_final.paragraphs)}")
    print(f" Jumlah Tabel: {len(doc_final.tables)}")
    print(f"=======================================================")

if __name__ == "__main__":
    assemble_full_skripsi()
