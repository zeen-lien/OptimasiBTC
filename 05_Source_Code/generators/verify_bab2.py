import docx

doc = docx.Document("03_Naskah_Skripsi/BAB_2_LANDASAN_TEORI_ZAINI_LEON.docx")

print("=" * 60)
print("AUDIT BAB 2 — FULL VERIFICATION")
print("=" * 60)

# 1. Check document structure: body children
tags = [c.tag.split('}')[-1] for c in doc._body._body]
print(f"\n1. Total body children: {len(tags)}")

# 2. Tables — should be 12: 1 matrix table + 11 equation tables
tbl_indices = [i for i, t in enumerate(tags) if t == 'tbl']
print(f"2. Total tables: {len(tbl_indices)}")
for idx in tbl_indices:
    print(f"   tbl at body index {idx}")

# 3. Check equations are NOT at the end (bug was: all 11 after sectPr)
sectPr_idx = tags.index('sectPr') if 'sectPr' in tags else len(tags)
eq_tables_after_sectPr = [i for i in tbl_indices if i > sectPr_idx]
print(f"\n3. sectPr at index: {sectPr_idx}")
print(f"   Equation tables AFTER sectPr (should be 0): {len(eq_tables_after_sectPr)}")
if len(eq_tables_after_sectPr) == 0:
    print("   ✅ FIX CONFIRMED: All tables are in proper sequence BEFORE sectPr!")
else:
    print("   ❌ BUG STILL PRESENT: Tables found after sectPr")

# 4. Check inline oMath elements in paragraphs (should be > 0 if math is inline)
inline_omath_count = 0
paras_with_inline_math = []
for i, p in enumerate(doc.paragraphs):
    omath_in_para = [c for c in p._element if c.tag.split('}')[-1] == 'oMath']
    if omath_in_para:
        inline_omath_count += len(omath_in_para)
        snippet = p.text[:80] if p.text else "(empty)"
        paras_with_inline_math.append((i, len(omath_in_para), snippet))

print(f"\n4. Inline oMath elements in paragraphs: {inline_omath_count}")
if inline_omath_count > 0:
    print("   ✅ FIX CONFIRMED: Inline math uses native OMML (no raw $...$ in text!)")
    for pi, cnt, snip in paras_with_inline_math:
        print(f"   Para {pi} has {cnt} oMath: \"{snip}...\"")
else:
    print("   ⚠️  No inline oMath found — check if $...$ tokens exist in source")

# 5. Check for raw $ signs in paragraph text (should be 0)
dollar_paras = []
for i, p in enumerate(doc.paragraphs):
    if '$' in p.text:
        dollar_paras.append((i, p.text[:80]))

print(f"\n5. Paragraphs with raw '$' signs (should be 0): {len(dollar_paras)}")
if len(dollar_paras) == 0:
    print("   ✅ FIX CONFIRMED: No raw LaTeX dollar signs in document text!")
else:
    print("   ❌ BUG STILL PRESENT: Raw $ found in:")
    for pi, snip in dollar_paras:
        print(f"   Para {pi}: \"{snip}...\"")

# 6. Total oMath in entire document (display + inline)
all_omath = doc._body._body.xpath(".//*[local-name()='oMath']")
print(f"\n6. Total oMath elements (display + inline): {len(all_omath)}")
print(f"   Expected: 11 display equations + {inline_omath_count} inline = {11 + inline_omath_count}")

# 7. Paper review headings
print("\n7. Paper Breakdown Headings:")
for p in doc.paragraphs:
    t = p.text.strip()
    if any(t.startswith(f"{i}. ") for i in range(1, 9)) and ("(" in t and ")" in t) and len(t) < 100:
        print(f"   ✔ {t}")

print("\n" + "=" * 60)
if len(eq_tables_after_sectPr) == 0 and len(dollar_paras) == 0 and inline_omath_count > 0:
    print("🎉 STATUS: KEDUA BUG BERHASIL DIPERBAIKI 100%!")
else:
    print("⚠️  STATUS: Masih ada yang perlu diperbaiki.")
print("=" * 60)
