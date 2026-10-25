"""
generate_front_matter.py
Membuat seluruh bagian awal (Front Matter) Skripsi UNISNU Jepara:
- Sampul Luar & Sampul Dalam
- Pernyataan Keaslian (Bebas Plagiasi)
- Persetujuan Pembimbing
- Pengesahan Tim Penguji
- Halaman Persembahan
- Halaman Motto
- Kata Pengantar
- Abstrak (Bahasa Indonesia) & Abstract (Bahasa Inggris)
- Daftar Isi
- Daftar Tabel
- Daftar Gambar
- Daftar Lampiran
"""

import os
import docx
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_font(run, name="Times New Roman", size=12, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic

def generate_front_matter(output_path="05_Source_Code/front_matter.docx"):
    doc = docx.Document()
    
    # Page Setup UNISNU: Top 4cm, Left 4cm, Bottom 3cm, Right 3cm
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(4.0)
    sec.left_margin = Cm(4.0)
    sec.bottom_margin = Cm(3.0)
    sec.right_margin = Cm(3.0)

    # Footer Roman Numeral
    sectPr = sec._sectPr
    pgNumType = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="lowerRoman" w:start="1"/>')
    sectPr.append(pgNumType)

    footer = sec.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f_run = f_p.add_run()
    set_font(f_run, "Times New Roman", 10)
    fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    f_p._p.append(fld)

    logo_path = os.path.abspath("05_Source_Code/extracted_logo/img_3.png")

    # =========================================================================
    # 1. SAMPUL LUAR (HARD COVER)
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run("OPTIMASI HYPERPARAMETER LONG SHORT-TERM MEMORY MENGGUNAKAN ALGORITMA OPTUNA PADA PREDIKSI HARGA BITCOIN")
    set_font(r, "Times New Roman", 14, bold=True)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(18)
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("SKRIPSI")
    set_font(r_sub, "Times New Roman", 14, bold=True)

    p_req = doc.add_paragraph()
    p_req.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_req.paragraph_format.space_before = Pt(6)
    p_req.paragraph_format.space_after = Pt(24)
    p_req.paragraph_format.line_spacing = 1.15
    r_req = p_req.add_run("Diajukan sebagai salah satu syarat untuk memperoleh gelar Sarjana Komputer (S.Kom.) pada Program Studi Teknik Informatika")
    set_font(r_req, "Times New Roman", 11, italic=False)

    # Logo
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(12)
        p_logo.paragraph_format.space_after = Pt(24)
        r_logo = p_logo.add_run()
        r_logo.add_picture(logo_path, width=Cm(5.0))

    p_auth = doc.add_paragraph()
    p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_auth.paragraph_format.space_before = Pt(12)
    p_auth.paragraph_format.space_after = Pt(4)
    p_auth.paragraph_format.line_spacing = 1.15
    r_auth = p_auth.add_run("Disusun Oleh:\nZAINI LEON MUSTHOFA KAMAL\nNIM: 231240001455")
    set_font(r_auth, "Times New Roman", 12, bold=True)

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(36)
    p_inst.paragraph_format.space_after = Pt(0)
    p_inst.paragraph_format.line_spacing = 1.15
    r_inst = p_inst.add_run("PROGRAM STUDI TEKNIK INFORMATIKA\nFAKULTAS SAINS DAN TEKNOLOGI\nUNIVERSITAS ISLAM NAHDLATUL ULAMA JEPARA\n2026")
    set_font(r_inst, "Times New Roman", 13, bold=True)

    doc.add_page_break()

    # =========================================================================
    # 2. SAMPUL DALAM
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run("OPTIMASI HYPERPARAMETER LONG SHORT-TERM MEMORY MENGGUNAKAN ALGORITMA OPTUNA PADA PREDIKSI HARGA BITCOIN")
    set_font(r, "Times New Roman", 14, bold=True)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(18)
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("SKRIPSI")
    set_font(r_sub, "Times New Roman", 14, bold=True)

    p_req = doc.add_paragraph()
    p_req.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_req.paragraph_format.space_before = Pt(6)
    p_req.paragraph_format.space_after = Pt(24)
    p_req.paragraph_format.line_spacing = 1.15
    r_req = p_req.add_run("Diajukan sebagai salah satu syarat untuk memperoleh gelar Sarjana Komputer (S.Kom.) pada Program Studi Teknik Informatika")
    set_font(r_req, "Times New Roman", 11, italic=False)

    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(12)
        p_logo.paragraph_format.space_after = Pt(24)
        r_logo = p_logo.add_run()
        r_logo.add_picture(logo_path, width=Cm(5.0))

    p_auth = doc.add_paragraph()
    p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_auth.paragraph_format.space_before = Pt(12)
    p_auth.paragraph_format.space_after = Pt(4)
    p_auth.paragraph_format.line_spacing = 1.15
    r_auth = p_auth.add_run("Disusun Oleh:\nZAINI LEON MUSTHOFA KAMAL\nNIM: 231240001455")
    set_font(r_auth, "Times New Roman", 12, bold=True)

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(36)
    p_inst.paragraph_format.space_after = Pt(0)
    p_inst.paragraph_format.line_spacing = 1.15
    r_inst = p_inst.add_run("PROGRAM STUDI TEKNIK INFORMATIKA\nFAKULTAS SAINS DAN TEKNOLOGI\nUNIVERSITAS ISLAM NAHDLATUL ULAMA JEPARA\n2026")
    set_font(r_inst, "Times New Roman", 13, bold=True)

    doc.add_page_break()

    # =========================================================================
    # 3. SURAT PERNYATAAN KEASLIAN
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("SURAT PERNYATAAN KEASLIAN KARYA TULIS SKRIPSI")
    set_font(r, "Times New Roman", 12, bold=True)

    p_stmt_intro = doc.add_paragraph()
    p_stmt_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_stmt_intro.paragraph_format.line_spacing = 1.5
    p_stmt_intro.paragraph_format.space_after = Pt(8)
    r = p_stmt_intro.add_run("Yang bertanda tangan di bawah ini:")
    set_font(r, "Times New Roman", 12)

    bio_tbl = doc.add_table(rows=4, cols=3)
    bio_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    bio_data = [
        ("Nama", ":", "Zaini Leon Musthofa Kamal"),
        ("NIM", ":", "231240001455"),
        ("Program Studi", ":", "Teknik Informatika"),
        ("Fakultas", ":", "Sains dan Teknologi, UNISNU Jepara")
    ]
    for idx, (lbl, col, val) in enumerate(bio_data):
        row = bio_tbl.rows[idx]
        cell_lbl, cell_col, cell_val = row.cells[0], row.cells[1], row.cells[2]
        cell_lbl.width = Cm(4.0)
        cell_col.width = Cm(0.5)
        cell_val.width = Cm(10.5)
        set_font(cell_lbl.paragraphs[0].add_run(lbl), "Times New Roman", 12)
        set_font(cell_col.paragraphs[0].add_run(col), "Times New Roman", 12)
        set_font(cell_val.paragraphs[0].add_run(val), "Times New Roman", 12, bold=(idx < 2))
        for c in (cell_lbl, cell_col, cell_val):
            c.paragraphs[0].paragraph_format.space_after = Pt(4)
            c.paragraphs[0].paragraph_format.line_spacing = 1.15

    p_stmt_body1 = doc.add_paragraph()
    p_stmt_body1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_stmt_body1.paragraph_format.line_spacing = 1.5
    p_stmt_body1.paragraph_format.first_line_indent = Cm(1.0)
    p_stmt_body1.paragraph_format.space_before = Pt(12)
    p_stmt_body1.paragraph_format.space_after = Pt(6)
    r = p_stmt_body1.add_run(
        "Menyatakan dengan sesungguhnya bahwa naskah Skripsi yang berjudul: \"Optimasi Hyperparameter Long Short-Term "
        "Memory Menggunakan Algoritma Optuna Pada Prediksi Harga Bitcoin\" adalah benar-benar hasil karya ilmiah mandiri "
        "saya sendiri, bukan merupakan karya tiruan, saduran, atau plagiasi dari karya orang lain, baik sebagian maupun seluruhnya. "
        "Sepanjang pengetahuan saya, tidak terdapat karya atau pendapat yang pernah ditulis atau diterbitkan oleh orang lain, "
        "kecuali yang secara tertulis dan formal disitasi dalam naskah ini serta dicantumkan dalam daftar pustaka sesuai dengan kaidah ilmiah."
    )
    set_font(r, "Times New Roman", 12)

    p_stmt_body2 = doc.add_paragraph()
    p_stmt_body2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_stmt_body2.paragraph_format.line_spacing = 1.5
    p_stmt_body2.paragraph_format.first_line_indent = Cm(1.0)
    p_stmt_body2.paragraph_format.space_after = Pt(18)
    r = p_stmt_body2.add_run(
        "Apabila di kemudian hari terbukti atau dapat dibuktikan bahwa skripsi ini merupakan hasil plagiasi, penjiplakan, atau rekayasa data "
        "yang melanggar kode etik akademik, maka saya bersedia menerima sanksi akademik yang seberat-beratnya sesuai dengan peraturan "
        "perundang-undangan yang berlaku dan Peraturan Rektor UNISNU Jepara, termasuk pencabutan gelar kesarjanaan yang telah saya peroleh."
    )
    set_font(r, "Times New Roman", 12)

    p_ttd = doc.add_paragraph()
    p_ttd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_ttd.paragraph_format.line_spacing = 1.15
    p_ttd.paragraph_format.space_before = Pt(18)
    r = p_ttd.add_run("Jepara, 20 Oktober 2026\nYang membuat pernyataan,\n\n\n\nMeterai Rp10.000\n\n\n")
    set_font(r, "Times New Roman", 12)
    r_nm = p_ttd.add_run("Zaini Leon Musthofa Kamal\nNIM: 231240001455")
    set_font(r_nm, "Times New Roman", 12, bold=True)

    doc.add_page_break()

    # =========================================================================
    # 4. LEMBAR PERSETUJUAN PEMBIMBING
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("LEMBAR PERSETUJUAN PEMBIMBING")
    set_font(r, "Times New Roman", 12, bold=True)

    p_app_intro = doc.add_paragraph()
    p_app_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_app_intro.paragraph_format.line_spacing = 1.5
    p_app_intro.paragraph_format.space_after = Pt(8)
    r = p_app_intro.add_run("Naskah Skripsi dengan identitas sebagai berikut:")
    set_font(r, "Times New Roman", 12)

    app_tbl = doc.add_table(rows=4, cols=3)
    app_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    app_data = [
        ("Nama Mahasiswa", ":", "Zaini Leon Musthofa Kamal"),
        ("Nomor Induk Mahasiswa", ":", "231240001455"),
        ("Program Studi", ":", "Teknik Informatika"),
        ("Judul Skripsi", ":", "Optimasi Hyperparameter Long Short-Term Memory Menggunakan Algoritma Optuna Pada Prediksi Harga Bitcoin")
    ]
    for idx, (lbl, col, val) in enumerate(app_data):
        row = app_tbl.rows[idx]
        cell_lbl, cell_col, cell_val = row.cells[0], row.cells[1], row.cells[2]
        cell_lbl.width = Cm(4.5)
        cell_col.width = Cm(0.5)
        cell_val.width = Cm(10.0)
        set_font(cell_lbl.paragraphs[0].add_run(lbl), "Times New Roman", 12)
        set_font(cell_col.paragraphs[0].add_run(col), "Times New Roman", 12)
        set_font(cell_val.paragraphs[0].add_run(val), "Times New Roman", 12, bold=(idx < 2))
        for c in (cell_lbl, cell_col, cell_val):
            c.paragraphs[0].paragraph_format.space_after = Pt(4)
            c.paragraphs[0].paragraph_format.line_spacing = 1.15

    p_app_body = doc.add_paragraph()
    p_app_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_app_body.paragraph_format.line_spacing = 1.5
    p_app_body.paragraph_format.first_line_indent = Cm(1.0)
    p_app_body.paragraph_format.space_before = Pt(14)
    p_app_body.paragraph_format.space_after = Pt(18)
    r = p_app_body.add_run(
        "Telah diperiksa dan disetujui oleh Dosen Pembimbing untuk diajukan dalam Ujian Sidang Skripsi (Munaqasyah) "
        "Program Studi Teknik Informatika Fakultas Sains dan Teknologi Universitas Islam Nahdlatul Ulama Jepara."
    )
    set_font(r, "Times New Roman", 12)

    p_app_ttd = doc.add_paragraph()
    p_app_ttd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_app_ttd.paragraph_format.line_spacing = 1.15
    p_app_ttd.paragraph_format.space_before = Pt(24)
    r = p_app_ttd.add_run("Jepara, 20 Oktober 2026\nMenyetujui,\nDosen Pembimbing,\n\n\n\n\n\n")
    set_font(r, "Times New Roman", 12)
    r_dos = p_app_ttd.add_run("Teguh Tamrin, S.Kom., M.Kom.\nNIDN. 0620127603 / NIY. 4 761220 17 235")
    set_font(r_dos, "Times New Roman", 12, bold=True)

    doc.add_page_break()

    # =========================================================================
    # 5. LEMBAR PENGESAHAN TIM PENGUJI
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("LEMBAR PENGESAHAN SKRIPSI")
    set_font(r, "Times New Roman", 12, bold=True)

    p_peng = doc.add_paragraph()
    p_peng.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_peng.paragraph_format.line_spacing = 1.5
    p_peng.paragraph_format.space_after = Pt(10)
    r = p_peng.add_run(
        "Skripsi dengan judul \"Optimasi Hyperparameter Long Short-Term Memory Menggunakan Algoritma Optuna Pada Prediksi Harga Bitcoin\" "
        "yang disusun oleh Zaini Leon Musthofa Kamal (NIM: 231240001455) telah dipertahankan di hadapan Dewan Penguji Ujian Sidang Skripsi "
        "Program Studi Teknik Informatika Fakultas Sains dan Teknologi Universitas Islam Nahdlatul Ulama Jepara pada tanggal 20 Oktober 2026 "
        "dan dinyatakan LULUS serta diterima sebagai salah satu syarat untuk memperoleh gelar Sarjana Komputer (S.Kom.)."
    )
    set_font(r, "Times New Roman", 12)

    p_dewan = doc.add_paragraph()
    p_dewan.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dewan.paragraph_format.space_before = Pt(10)
    p_dewan.paragraph_format.space_after = Pt(8)
    r = p_dewan.add_run("DEWAN PENGUJI:")
    set_font(r, "Times New Roman", 12, bold=True)

    t_peng = doc.add_table(rows=3, cols=3)
    t_peng.alignment = WD_TABLE_ALIGNMENT.CENTER
    peng_data = [
        ("Ketua Penguji", ":", "Teguh Tamrin, S.Kom., M.Kom."),
        ("Sekretaris Penguji", ":", "Gentur Wahyu Nyipto Wibowo, S.Kom., M.Kom."),
        ("Anggota Penguji", ":", "Danang Mahendra, S.Kom., M.Kom.")
    ]
    for idx, (lbl, col, val) in enumerate(peng_data):
        row = t_peng.rows[idx]
        row.cells[0].width = Cm(4.5)
        row.cells[1].width = Cm(0.5)
        row.cells[2].width = Cm(10.0)
        set_font(row.cells[0].paragraphs[0].add_run(lbl), "Times New Roman", 12)
        set_font(row.cells[1].paragraphs[0].add_run(col), "Times New Roman", 12)
        set_font(row.cells[2].paragraphs[0].add_run(val + "  ( ..................................... )"), "Times New Roman", 12)
        for c in row.cells:
            c.paragraphs[0].paragraph_format.space_after = Pt(8)
            c.paragraphs[0].paragraph_format.line_spacing = 1.15

    p_mengk = doc.add_paragraph()
    p_mengk.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mengk.paragraph_format.line_spacing = 1.15
    p_mengk.paragraph_format.space_before = Pt(20)
    r = p_mengk.add_run("Mengesahkan,\nDekan Fakultas Sains dan Teknologi\nUniversitas Islam Nahdlatul Ulama Jepara\n\n\n\n\n")
    set_font(r, "Times New Roman", 12)
    r_dek = p_mengk.add_run("Dr. Arif Mustofa, S.T., M.Si.\nNIY. 3 730617 13 124 / NIDN. 0617067302")
    set_font(r_dek, "Times New Roman", 12, bold=True)

    doc.add_page_break()

    # =========================================================================
    # 6. HALAMAN PERSEMBAHAN
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("HALAMAN PERSEMBAHAN")
    set_font(r, "Times New Roman", 12, bold=True)

    p_persembahan = doc.add_paragraph()
    p_persembahan.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_persembahan.paragraph_format.line_spacing = 1.5
    p_persembahan.paragraph_format.space_after = Pt(12)
    r = p_persembahan.add_run(
        "Alhamdulillahirabbil'alamin, dengan penuh rasa syukur ke hadirat Allah SWT atas segala limpahan rahmat, hidayah, "
        "dan kemudahan yang dianugerahkan-Nya, karya tulis skripsi yang sederhana ini saya persembahkan dengan segenap cinta, "
        "penghargaan, dan rasa hormat yang mendalam kepada:"
    )
    set_font(r, "Times New Roman", 12)

    persembahan_items = [
        "Kedua orang tua tercinta yang tiada henti memanjatkan doa tulus di setiap sujudnya, memberikan dukungan moril dan materil tanpa pamrih, serta menanamkan nilai-nilai kegigihan dan ketulusan dalam menuntut ilmu.",
        "Keluarga besar yang senantiasa menjadi sumber inspirasi, motivasi, dan kehangatan dalam setiap langkah perjuangan menempuh studi pendidikan sarjana.",
        "Dosen Pembimbing dan seluruh Bapak/Ibu Dosen Program Studi Teknik Informatika FST UNISNU Jepara yang telah ikhlas mentransfer ilmu pengetahuan, bimbingan, arahan metodologis, dan teladan integritas akademik yang tak ternilai.",
        "Rekan-rekan seperjuangan mahasiswa Program Studi Teknik Informatika FST UNISNU Jepara Angkatan 2023 atas kebersamaan, sinergi, diskusi intelektual, dan persahabatan yang senantiasa saling menguatkan.",
        "Almamater tercinta, Universitas Islam Nahdlatul Ulama (UNISNU) Jepara, sebagai kawah candradimuka yang menempa keilmuan teknologi informasi berlandaskan nilai-nilai Ahlussunnah wal Jama'ah."
    ]
    for item in persembahan_items:
        p_item = doc.add_paragraph()
        p_item.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_item.paragraph_format.line_spacing = 1.5
        p_item.paragraph_format.first_line_indent = Cm(1.0)
        p_item.paragraph_format.space_after = Pt(6)
        r = p_item.add_run("•  " + item)
        set_font(r, "Times New Roman", 12)

    doc.add_page_break()

    # =========================================================================
    # 7. HALAMAN MOTTO
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("HALAMAN MOTTO")
    set_font(r, "Times New Roman", 12, bold=True)

    mottos = [
        ("“Barangsiapa bertakwa kepada Allah niscaya Dia akan mengadakan baginya jalan keluar, dan memberinya rezeki dari arah yang tiada disangka-sangkanya. Dan barangsiapa yang bertawakkal kepada Allah niscaya Allah akan mencukupkan (keperluan)nya.”", "(QS. At-Talaq: 2-3)"),
        ("“Sesungguhnya Allah mencintai seorang hamba yang apabila mengerjakan suatu pekerjaan, ia melakukannya dengan itqan (profesional, tekun, dan tuntas).”", "(HR. At-Thabrani dan Al-Baihaqi)"),
        ("“Bukan kecerdasan semata yang menentukan keberhasilan penyelesaian suatu riset, melainkan ketekunan dalam menghadapi anomali data, kedisiplinan metode ilmiah, dan kerendahan hati untuk terus belajar memperbarui pemahaman.”", "(Zaini Leon Musthofa Kamal)")
    ]

    for quote, src in mottos:
        p_q = doc.add_paragraph()
        p_q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_q.paragraph_format.line_spacing = 1.5
        p_q.paragraph_format.space_before = Pt(18)
        p_q.paragraph_format.space_after = Pt(4)
        r_q = p_q.add_run(quote)
        set_font(r_q, "Times New Roman", 12, italic=True)

        p_s = doc.add_paragraph()
        p_s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_s.paragraph_format.space_after = Pt(24)
        r_s = p_s.add_run(src)
        set_font(r_s, "Times New Roman", 11, bold=True)

    doc.add_page_break()

    # =========================================================================
    # 8. KATA PENGANTAR
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("KATA PENGANTAR")
    set_font(r, "Times New Roman", 12, bold=True)

    kata_pengantar_paras = [
        "Puji dan syukur senantiasa penulis panjatkan ke hadirat Allah SWT atas limpahan rahmat, taufiq, karunia, dan inayah-Nya, sehingga penulis dapat menyelesaikan penyusunan skripsi yang berjudul \"Optimasi Hyperparameter Long Short-Term Memory Menggunakan Algoritma Optuna Pada Prediksi Harga Bitcoin\" dengan baik dan lancar. Shalawat serta salam semoga senantiasa tercurahkan kepada junjungan kita Nabi Agung Muhammad SAW, keluarga, para sahabat, dan pengikutnya hingga akhir zaman.",
        "Skripsi ini disusun sebagai salah satu syarat akademik guna menyelesaikan studi jenjang Strata 1 (S1) dan memperoleh gelar Sarjana Komputer (S.Kom.) pada Program Studi Teknik Informatika, Fakultas Sains dan Teknologi, Universitas Islam Nahdlatul Ulama (UNISNU) Jepara. Penulis menyadari sepenuhnya bahwa selesainya penulisan skripsi ini tidak terlepas dari bantuan, bimbingan, doa, dan motivasi dari berbagai pihak.",
        "Oleh karena itu, dengan penuh kerendahan hati dan rasa hormat yang mendalam, penulis ingin menyampaikan ucapan terima kasih yang sebesar-besarnya kepada:",
        "1. Prof. Dr. H. Abdul Djamil, M.A., selaku Rektor Universitas Islam Nahdlatul Ulama (UNISNU) Jepara, atas segala fasilitas dan kebijakan akademik yang mendukung kelancaran studi mahasiswa.\n"
        "2. Dr. Arif Mustofa, S.T., M.Si., selaku Dekan Fakultas Sains dan Teknologi UNISNU Jepara, atas kepemimpinan, dedikasi, dan sarana prasarana yang disediakan selama masa perkuliahan.\n"
        "3. Gentur Wahyu Nyipto Wibowo, S.Kom., M.Kom., selaku Ketua Program Studi Teknik Informatika FST UNISNU Jepara, atas pengarahan kurikulum, dorongan akademis, dan kemudahan proses administrasi tugas akhir.\n"
        "4. Teguh Tamrin, S.Kom., M.Kom., selaku Dosen Pembimbing Skripsi, yang telah dengan penuh kesabaran, ketelitian, dan keikhlasan meluangkan waktu serta mencurahkan pemikiran dalam membimbing, mengarahkan, dan memberikan telaah kritis sejak awal pengajuan proposal hingga naskah skripsi ini terselesaikan secara utuh dan teruji.\n"
        "5. Seluruh Dewan Penguji Ujian Sidang Skripsi yang telah memberikan koreksi, evaluasi metodologis, dan masukan berharga demi penyempurnaan kualitas ilmiah naskah ini.\n"
        "6. Seluruh Bapak dan Ibu Dosen serta staf akademik di lingkungan Program Studi Teknik Informatika dan Fakultas Sains dan Teknologi UNISNU Jepara yang telah membekali penulis dengan berbagai khazanah ilmu pengetahuan dan keteladanan budi pekerti luhur.\n"
        "7. Kedua orang tua tercinta dan segenap keluarga besar yang senantiasa mengiringi setiap perjuangan penulis dengan doa yang tulus, limpahan kasih sayang yang tak bertepi, serta pengorbanan moril dan materiil yang tak terhingga nilainya.\n"
        "8. Sahabat-sahabat dan rekan mahasiswa Program Studi Teknik Informatika UNISNU Jepara Angkatan 2023 atas kebersamaan, kekompakan, serta diskusi saintifik yang selalu membangun.",
        "Penulis menyadari bahwa skripsi ini masih jauh dari kesempurnaan akibat keterbatasan pengalaman dan wawasan penulis. Meskipun demikian, penulis telah berusaha semaksimal mungkin mencurahkan kemampuan terbaik dalam menyelesaikan karya ini. Akhir kata, penulis berharap semoga skripsi ini dapat memberikan kontribusi nyata serta kemaslahatan bagi pengembangan ilmu pengetahuan dan teknologi, khususnya di bidang Artificial Intelligence, Deep Learning, dan rekayasa optimasi peramalan finansial."
    ]

    for idx, text in enumerate(kata_pengantar_paras):
        p_kp = doc.add_paragraph()
        p_kp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_kp.paragraph_format.line_spacing = 1.5
        p_kp.paragraph_format.space_after = Pt(6)
        if idx not in [0, 2]:
            p_kp.paragraph_format.first_line_indent = Cm(1.0)
        r = p_kp.add_run(text)
        set_font(r, "Times New Roman", 12)

    p_kp_ttd = doc.add_paragraph()
    p_kp_ttd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_kp_ttd.paragraph_format.line_spacing = 1.15
    p_kp_ttd.paragraph_format.space_before = Pt(16)
    r = p_kp_ttd.add_run("Jepara, 20 Oktober 2026\nPenulis,\n\n\n\n")
    set_font(r, "Times New Roman", 12)
    r_nm = p_kp_ttd.add_run("Zaini Leon Musthofa Kamal\nNIM: 231240001455")
    set_font(r_nm, "Times New Roman", 12, bold=True)

    doc.add_page_break()

    # =========================================================================
    # 9. ABSTRAK (BAHASA INDONESIA)
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("ABSTRAK")
    set_font(r, "Times New Roman", 12, bold=True)

    p_id = doc.add_paragraph()
    p_id.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_id.paragraph_format.space_after = Pt(10)
    p_id.paragraph_format.line_spacing = 1.0
    r_id = p_id.add_run("Zaini Leon Musthofa Kamal, 231240001455, Optimasi Hyperparameter Long Short-Term Memory Menggunakan Algoritma Optuna Pada Prediksi Harga Bitcoin, 2026.")
    set_font(r_id, "Times New Roman", 11, bold=True)

    # 1 Single Paragraph, 250 - 500 words
    abstrak_id_text = (
        "Volatilitas ekstrem dan sifat non-stasioneritas pada harga penutupan Bitcoin memicu tantangan signifikan dalam "
        "pemodelan deret waktu finansial. Arsitektur Long Short-Term Memory (LSTM) terbukti andal dalam menangkap ketergantungan "
        "jangka panjang, namun kinerjanya sangat sensitif terhadap kombinasi hyperparameter arsitektural dan pelatihan. Pendekatan "
        "penalaan manual maupun pencarian konvensional rentan terjebak pada local minima dan memerlukan biaya komputasi yang tinggi. "
        "Penelitian ini bertujuan untuk merancang kerangka kerja optimasi otomatis berbasis Bayesian Optimization dengan algoritma "
        "Tree-structured Parzen Estimator (TPE) melalui framework Optuna guna menemukan konfigurasi hyperparameter LSTM yang optimal, "
        "serta membuktikan signifikansi penurunan error secara empiris dibandingkan konfigurasi default dan manual tuning. Eksperimen "
        "menggunakan 2.192 data transaksi harian BTC/USDT dari Binance Public API periode 1 Januari 2020 hingga 31 Desember 2025. "
        "Data dipartisi secara kronologis menjadi 80% data latih in-sample (1.753 observasi) dan 20% data uji out-of-sample holdout "
        "(439 observasi) dengan normalisasi Min-Max skalar tunggal untuk mencegah kebocoran temporal. Ruang pencarian Optuna-TPE "
        "mencakup jendela historis (lookback window), jumlah unit tersembunyi, laju dropout, learning rate, dan ukuran batch dengan "
        "skema validasi Time-Series Split 5-fold serta mekanisme pemangkasan MedianPruner sebanyak 50 trials. Evaluasi model "
        "dilakukan menggunakan tiga metrik utama: Root Mean Square Error (RMSE), Mean Absolute Error (MAE), dan Mean Absolute Percentage "
        "Error (MAPE). Hasil pengujian empiris membuktikan bahwa arsitektur usulan Tier 3 (Optuna-TPE dengan MAE Loss dan ReduceLROnPlateau) berhasil mencapai performa "
        "unggul dengan nilai RMSE USD 2.338,22, MAE USD 1.784,27, dan MAPE 1,81%. Model ini melampaui model rujukan Tier 2 "
        "(Manual Tuning Khumaidi dkk.) dengan penurunan galat sangat signifikan mencapai 40,11% pada RMSE (USD 3.904,15), 43,94% pada MAE (USD 3.182,86), "
        "serta 42,72% pada MAPE (3,16%), sekaligus mengungguli model baseline default Tier 1 (RMSE USD 2.947,60; MAPE 2,36% dengan penurunan 20,67% RMSE dan 23,31% MAPE). Konfigurasi hyperparameter optimal yang ditemukan terdiri atas lookback window 22 hari, 100 hidden units, "
        "dropout 0,10, laju pembelajaran adaptif 0,00197 dengan scheduler dynamic decay, serta batch size 16. Temuan ini membuktikan bahwa penalaan otomatis berbasis Optuna-TPE terarah "
        "(domain-guided) yang dipadukan dengan fungsi kerugian linier MAE secara efektif mengatasi keterlambatan momentum volatilitas dan menghasilkan generalisasi peramalan tingkat presisi tinggi di bawah 2% pada data deret waktu aset kripto."
    )
    p_abs = doc.add_paragraph()
    p_abs.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_abs.paragraph_format.first_line_indent = Cm(1.0)
    p_abs.paragraph_format.line_spacing = 1.0
    p_abs.paragraph_format.space_after = Pt(12)
    r = p_abs.add_run(abstrak_id_text)
    set_font(r, "Times New Roman", 11)

    p_kw = doc.add_paragraph()
    p_kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kw.paragraph_format.line_spacing = 1.0
    p_kw.paragraph_format.space_after = Pt(12)
    r_kw_lbl = p_kw.add_run("Kata Kunci: ")
    set_font(r_kw_lbl, "Times New Roman", 11, bold=True)
    r_kw_val = p_kw.add_run("Long Short-Term Memory, Optuna-TPE, Prediksi Harga Bitcoin, Optimasi Hyperparameter, Deret Waktu Finansial.")
    set_font(r_kw_val, "Times New Roman", 11, italic=True)

    doc.add_page_break()

    # =========================================================================
    # 10. ABSTRACT (BAHASA INGGRIS)
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("ABSTRACT")
    set_font(r, "Times New Roman", 12, bold=True)

    p_en_id = doc.add_paragraph()
    p_en_id.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_en_id.paragraph_format.space_after = Pt(10)
    p_en_id.paragraph_format.line_spacing = 1.0
    r_en_id = p_en_id.add_run("Zaini Leon Musthofa Kamal, 231240001455, Hyperparameter Optimization of Long Short-Term Memory Using Optuna Algorithm in Bitcoin Price Prediction, 2026.")
    set_font(r_en_id, "Times New Roman", 11, bold=True)

    abstrak_en_text = (
        "Extreme volatility and non-stationarity in Bitcoin closing prices pose substantial challenges to financial "
        "time-series modeling. While Long Short-Term Memory (LSTM) networks excel at learning long-term dependencies, their predictive "
        "performance remains highly sensitive to hyperparameter configurations. Traditional trial-and-error adjustments and exhaustive "
        "search methods frequently suffer from prohibitive computational costs and risk converging to suboptimal local minima. This study "
        "proposes an automated hyperparameter optimization framework utilizing the Tree-structured Parzen Estimator (TPE) algorithm within "
        "the Optuna framework to systematically identify optimal LSTM configurations, establishing empirical performance gains over baseline "
        "and manual tuning benchmarks. The empirical dataset comprises 2,192 daily BTC/USDT transaction records retrieved via the Binance "
        "Public API spanning January 1, 2020, to December 31, 2025. Data were chronologically partitioned into an 80% in-sample training "
        "subset (1,753 observations) and a 20% out-of-sample holdout test subset (439 observations), normalized using in-sample single-scale "
        "Min-Max scaling to prevent temporal data leakage. The search space encompassed lookback windows, hidden units, dropout rates, learning "
        "rates, and batch sizes, evaluated across 50 trials via a 5-fold Time-Series Split and MedianPruner mechanism, integrated with "
        "MAE loss training and dynamic learning rate decay. Model performance was assessed using Root Mean Square Error (RMSE), Mean Absolute "
        "Error (MAE), and Mean Absolute Percentage Error (MAPE). Empirical findings demonstrate that the proposed Tier 3 (Optuna-TPE) "
        "architecture achieved superior forecasting accuracy across all models, yielding an RMSE of USD 2,338.22, MAE of USD 1,784.27, "
        "and MAPE of 1.81%. This represents a 40.11% error reduction in RMSE, 43.94% in MAE, and 42.72% in MAPE over Tier 2 Manual Tuning "
        "(RMSE USD 3,904.15; MAE USD 3,182.86; MAPE 3.16%), while outperforming the Tier 1 baseline default (RMSE USD 2,947.60; MAPE 2.36% with a 20.67% RMSE reduction). "
        "The optimal configuration comprised a 22-day lookback window, 100 hidden units, 0.10 dropout, 0.00197 learning rate, and batch size 16. "
        "These results confirm that Bayesian TPE optimization with domain-guided bounds and loss-metric alignment effectively mitigates momentum lag "
        "and enhances generalization capability in cryptocurrency predictive analytics."
    )
    p_abs_en = doc.add_paragraph()
    p_abs_en.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_abs_en.paragraph_format.first_line_indent = Cm(1.0)
    p_abs_en.paragraph_format.line_spacing = 1.0
    p_abs_en.paragraph_format.space_after = Pt(12)
    r = p_abs_en.add_run(abstrak_en_text)
    set_font(r, "Times New Roman", 11, italic=True)

    p_kw_en = doc.add_paragraph()
    p_kw_en.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kw_en.paragraph_format.line_spacing = 1.0
    p_kw_en.paragraph_format.space_after = Pt(12)
    r_kw_lbl_en = p_kw_en.add_run("Keywords: ")
    set_font(r_kw_lbl_en, "Times New Roman", 11, bold=True)
    r_kw_val_en = p_kw_en.add_run("Long Short-Term Memory, Optuna-TPE, Bitcoin Price Forecasting, Hyperparameter Optimization, Financial Time-Series.")
    set_font(r_kw_val_en, "Times New Roman", 11, italic=True)

    doc.add_page_break()

    # =========================================================================
    # 11. DAFTAR ISI
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("DAFTAR ISI")
    set_font(r, "Times New Roman", 12, bold=True)

    toc_entries = [
        ("HALAMAN SAMPUL LUAR", "i", False),
        ("HALAMAN SAMPUL DALAM", "ii", False),
        ("SURAT PERNYATAAN KEASLIAN", "iii", False),
        ("LEMBAR PERSETUJUAN PEMBIMBING", "iv", False),
        ("LEMBAR PENGESAHAN SKRIPSI", "v", False),
        ("HALAMAN PERSEMBAHAN", "vi", False),
        ("HALAMAN MOTTO", "vii", False),
        ("KATA PENGANTAR", "viii", False),
        ("ABSTRAK", "ix", False),
        ("ABSTRACT", "x", False),
        ("DAFTAR ISI", "xi", False),
        ("DAFTAR TABEL", "xii", False),
        ("DAFTAR GAMBAR", "xiii", False),
        ("DAFTAR LAMPIRAN", "xiv", False),
        ("BAB I   PENDAHULUAN", "1", True),
        ("   A. Latar Belakang Masalah", "1", False),
        ("   B. Rumusan Masalah", "6", False),
        ("   C. Batasan Masalah", "7", False),
        ("   D. Tujuan Penelitian", "8", False),
        ("   E. Manfaat Penelitian", "9", False),
        ("   F. Keaslian Penelitian", "10", False),
        ("   G. Sistematika Penulisan", "11", False),
        ("BAB II  LANDASAN TEORI", "13", True),
        ("   A. Penelitian Terdahulu", "13", False),
        ("   B. Landasan Teori", "20", False),
        ("      1. Karakteristik Data Deret Waktu Finansial Bitcoin", "20", False),
        ("      2. Normalisasi Data Min-Max Scaling", "22", False),
        ("      3. Arsitektur Jaringan Saraf Tiruan LSTM", "24", False),
        ("      4. Regularisasi Dropout", "27", False),
        ("      5. Algoritma Optimasi Bobot Adam", "29", False),
        ("      6. Algoritma Optuna dan Tree-structured Parzen Estimator", "31", False),
        ("      7. Mekanisme Pemangkasan MedianPruner", "34", False),
        ("      8. Metrik Evaluasi Kinerja Peramalan (RMSE, MAE, MAPE)", "36", False),
        ("   C. Kerangka Berpikir", "39", False),
        ("BAB III METODE PENELITIAN", "42", True),
        ("   A. Jenis dan Pendekatan Penelitian", "42", False),
        ("   B. Tempat dan Waktu Penelitian", "44", False),
        ("   C. Sumber Data dan Variabel Penelitian", "46", False),
        ("   D. Metode Pengumpulan Data", "49", False),
        ("   E. Kerangka Kerja dan Tahapan Penelitian", "51", False),
        ("   F. Metode Analisis dan Evaluasi Model", "56", False),
        ("BAB IV  HASIL PENELITIAN DAN PEMBAHASAN", "60", True),
        ("   A. Hasil Penelitian", "60", False),
        ("      1. Hasil Pengumpulan dan Audit Kualitas Data Historis Binance", "60", False),
        ("      2. Hasil Pra-Pemrosesan Data dan Partisi Deret Waktu", "63", False),
        ("      3. Hasil Pelatihan Model Baseline Default (Tier 1)", "65", False),
        ("      4. Hasil Pelatihan Model Manual Tuning (Tier 2)", "67", False),
        ("      5. Hasil Optimasi Hyperparameter Optuna-TPE (Tier 3)", "69", False),
        ("      6. Perbandingan Hasil Evaluasi Kinerja Out-of-Sample", "72", False),
        ("   B. Pembahasan", "74", False),
        ("      1. Analisis Efektivitas Algoritma Optuna-TPE", "74", False),
        ("      2. Fenomena Overfitting pada Manual Tuning", "76", False),
        ("      3. Pengaruh Jendela Waktu Pendek Terhadap Volatilitas", "78", False),
        ("      4. Implikasi Praktis dan Kontribusi Penelitian", "80", False),
        ("BAB V   PENUTUP", "82", True),
        ("   A. Simpulan", "82", False),
        ("   B. Saran", "84", False),
        ("DAFTAR PUSTAKA", "86", True),
        ("LAMPIRAN", "91", True),
        ("   Lampiran 1: Ringkasan Kode Program Pipeline Komputasi Python", "91", False),
        ("   Lampiran 2: Sampel Data Historis Transaksi Binance BTC/USDT", "96", False),
        ("   Lampiran 3: Log Riwayat Percobaan 50 Trial Optuna-TPE", "100", False),
    ]

    for title, page_no, is_bold in toc_entries:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.line_spacing = 1.15
        p_t.paragraph_format.space_after = Pt(2)
        p_t.paragraph_format.tab_stops.add_tab_stop(Cm(13.8), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r_t = p_t.add_run(title)
        set_font(r_t, "Times New Roman", 11, bold=is_bold)
        r_tab = p_t.add_run(f"\t{page_no}")
        set_font(r_tab, "Times New Roman", 11, bold=is_bold)

    doc.add_page_break()

    # =========================================================================
    # 12. DAFTAR TABEL
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("DAFTAR TABEL")
    set_font(r, "Times New Roman", 12, bold=True)

    table_entries = [
        ("Tabel 2.1", "Matriks Sintesis Literatur Komparatif Penelitian Terdahulu", "19"),
        ("Tabel 3.1", "Jadwal Pelaksanaan Tahapan Penelitian Tahun 2026", "45"),
        ("Tabel 3.2", "Deskripsi Fitur Data Historis Binance BTC/USDT", "48"),
        ("Tabel 3.3", "Ruang Pencarian (Search Space) Hyperparameter Optuna-TPE", "54"),
        ("Tabel 3.4", "Spesifikasi Perangkat Keras dan Lingkungan Komputasi Eksperimen", "55"),
        ("Tabel 3.5", "Skala Penilaian Interpretasi Nilai Mean Absolute Percentage Error (MAPE)", "58"),
        ("Tabel 4.1", "Statistik Deskriptif Data Historis Harian BTC/USDT (2020–2025)", "61"),
        ("Tabel 4.2", "Ringkasan Hasil Partisi Data dan Parameter Skalasi Min-Max", "64"),
        ("Tabel 4.3", "Top 5 Konfigurasi Hyperparameter Terbaik Hasil Eksperimen Optuna-TPE", "70"),
        ("Tabel 4.4", "Evaluasi Kinerja Komparatif Model 3-Tier pada Out-of-Sample Test Set", "72"),
        ("Tabel 4.5", "Persentase Penurunan Nilai Metrik Error (Positive Improvement)", "74"),
    ]

    for tbl_num, tbl_title, page_no in table_entries:
        p_tb = doc.add_paragraph()
        p_tb.paragraph_format.line_spacing = 1.15
        p_tb.paragraph_format.space_after = Pt(3)
        p_tb.paragraph_format.tab_stops.add_tab_stop(Cm(13.8), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r_num = p_tb.add_run(f"{tbl_num}  {tbl_title}")
        set_font(r_num, "Times New Roman", 11, bold=False)
        r_p = p_tb.add_run(f"\t{page_no}")
        set_font(r_p, "Times New Roman", 11, bold=False)

    doc.add_page_break()

    # =========================================================================
    # 13. DAFTAR GAMBAR
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("DAFTAR GAMBAR")
    set_font(r, "Times New Roman", 12, bold=True)

    figure_entries = [
        ("Gambar 2.1", "Diagram Alir Kerangka Berpikir Optimasi Hyperparameter LSTM", "41"),
        ("Gambar 3.1", "Alur Metodologi Eksperimen Komputasi Tiga Tingkat (3-Tier)", "52"),
        ("Gambar 4.1", "Grafik Perbandingan Nilai Aktual vs Prediksi Model 3-Tier (439 Hari Out-of-Sample)", "73"),
        ("Gambar 4.2", "Kurva Konvergensi Optimasi Bayesian Optuna-TPE Sepanjang 50 Trials", "75"),
        ("Gambar 4.3", "Distribusi Densitas Residual Error Prediksi Out-of-Sample (Tier 1 vs Tier 2 vs Tier 3)", "77"),
    ]

    for fig_num, fig_title, page_no in figure_entries:
        p_fg = doc.add_paragraph()
        p_fg.paragraph_format.line_spacing = 1.15
        p_fg.paragraph_format.space_after = Pt(3)
        p_fg.paragraph_format.tab_stops.add_tab_stop(Cm(13.8), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r_num = p_fg.add_run(f"{fig_num}  {fig_title}")
        set_font(r_num, "Times New Roman", 11, bold=False)
        r_p = p_fg.add_run(f"\t{page_no}")
        set_font(r_p, "Times New Roman", 11, bold=False)

    doc.add_page_break()

    # =========================================================================
    # 14. DAFTAR LAMPIRAN
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("DAFTAR LAMPIRAN")
    set_font(r, "Times New Roman", 12, bold=True)

    appendix_entries = [
        ("Lampiran 1", "Ringkasan Kode Program Pipeline Komputasi Python (01 s.d. 06)", "91"),
        ("Lampiran 2", "Sampel Data Historis Transaksi Binance BTC/USDT Periode 2020–2025", "96"),
        ("Lampiran 3", "Log Riwayat Percobaan 50 Trial Optimasi Bayesian Optuna-TPE", "100"),
    ]

    for app_num, app_title, page_no in appendix_entries:
        p_ap = doc.add_paragraph()
        p_ap.paragraph_format.line_spacing = 1.15
        p_ap.paragraph_format.space_after = Pt(3)
        p_ap.paragraph_format.tab_stops.add_tab_stop(Cm(13.8), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r_num = p_ap.add_run(f"{app_num}: {app_title}")
        set_font(r_num, "Times New Roman", 11, bold=False)
        r_p = p_ap.add_run(f"\t{page_no}")
        set_font(r_p, "Times New Roman", 11, bold=False)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Front matter generated successfully: {output_path} ({os.path.getsize(output_path)} bytes)")

if __name__ == "__main__":
    generate_front_matter()
